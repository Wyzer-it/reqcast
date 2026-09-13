import docx

from reqcast.docx_extract import extract_docx

_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d494844520000000a0000000a08020000"
    "00025058ea0000001649444154789c63fcffff3f036ec084478e61e4"
    "4a0300a5e30311c77a1c550000000049454e44ae426082"
)


def test_docx_heading_styles_and_labels(tmp_path, base_cfg):
    d = docx.Document()
    d.add_heading("1.1 Section", level=2)
    d.add_paragraph("REQ-001 First")
    d.add_paragraph("Description: The system shall do the first thing.")
    path = tmp_path / "spec.docx"
    d.save(str(path))

    roots, report = extract_docx(path, base_cfg)
    assert report.headings == 1
    assert report.requirements == 1
    heading = roots[0]
    assert heading.kind == "heading"
    req = heading.children[0]
    assert req.identifier == "REQ-001"
    assert "shall do the first thing" in req.body


def test_docx_image_and_caption_pair_below(tmp_path, base_cfg):
    d = docx.Document()
    d.add_paragraph("REQ-001 First")
    d.add_paragraph("Description: text.")
    img_path = tmp_path / "fig.png"
    img_path.write_bytes(_PNG)
    d.add_picture(str(img_path))
    d.add_paragraph("Figure 1: System overview")
    d.save(str(tmp_path / "spec.docx"))

    roots, report = extract_docx(tmp_path / "spec.docx", base_cfg)
    assert report.figures_found == 1
    assert report.figures_cropped == 1
    assert roots[0].figures[0].caption == "Figure 1: System overview"
    assert roots[0].figures[0].image_png == _PNG


def test_docx_orphan_image_does_not_leak_into_next_section(tmp_path, base_cfg):
    """Regression: an image with no caption before a new heading used to stay
    in `pending_images` forever and get attached to a LATER, unrelated
    captioned figure in the next section instead of being reported as
    orphaned."""
    d = docx.Document()
    d.add_paragraph("REQ-001 First")
    d.add_paragraph("Description: text one.")
    img_path = tmp_path / "orphan.png"
    img_path.write_bytes(_PNG)
    d.add_picture(str(img_path))  # no caption follows - a logo, say

    d.add_heading("1.2 Next section", level=2)
    d.add_paragraph("REQ-002 Second")
    d.add_paragraph("Description: text two.")
    img_path2 = tmp_path / "real.png"
    img_path2.write_bytes(_PNG)
    d.add_picture(str(img_path2))
    d.add_paragraph("Figure 1: The real figure")
    d.save(str(tmp_path / "spec.docx"))

    roots, report = extract_docx(tmp_path / "spec.docx", base_cfg)
    assert report.figures_found == 2
    assert report.figures_cropped == 1
    assert len(report.figures_skipped) == 1

    req1 = roots[0]
    assert req1.identifier == "REQ-001"
    assert req1.figures == []  # the orphan must NOT have attached here

    heading = roots[1]
    req2 = heading.children[0]
    assert req2.identifier == "REQ-002"
    assert len(req2.figures) == 1
    assert req2.figures[0].caption == "Figure 1: The real figure"


def test_docx_orphan_caption_with_no_image_is_reported(tmp_path, base_cfg):
    d = docx.Document()
    d.add_paragraph("REQ-001 First")
    d.add_paragraph("Description: text.")
    d.add_paragraph("Figure 1: Never shows up")  # caption with no image at all
    d.add_heading("1.2 Next", level=2)
    d.add_paragraph("REQ-002 Second")
    d.add_paragraph("Description: text two.")
    d.save(str(tmp_path / "spec.docx"))

    roots, report = extract_docx(tmp_path / "spec.docx", base_cfg)
    assert report.figures_found == 1
    assert report.figures_cropped == 0
    assert len(report.figures_skipped) == 1
    assert roots[0].figures == []
