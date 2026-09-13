"""Builds examples/demo-spec.docx: the same fabricated "Aria Smart
Thermostat" example as demo-spec.pdf/.md/.txt, but using real Word
Heading styles and a real embedded image - to demonstrate the .docx
extractor's format-native path (docx_extract.py) rather than forcing it
through PDF-style pattern matching.

Not part of the reqcast package - a one-off generator, kept here so the
example can be regenerated without hand-editing a binary .docx.
Run: .venv/bin/python examples/_build_demo_spec_docx.py
"""
from __future__ import annotations

import pathlib

import docx
from docx.shared import Inches

HERE = pathlib.Path(__file__).parent
OUT = HERE / "demo-spec.docx"


def build() -> None:
    d = docx.Document()
    d.add_heading("Aria Smart Thermostat", level=1)
    d.add_paragraph("Software Requirements Specification (fabricated example document)")

    d.add_heading("3.1 Power Management", level=2)
    d.add_paragraph("REQ-DEMO-001 Low-power standby")
    d.add_paragraph(
        "Description: The thermostat shall enter standby mode after 120 seconds of no "
        "user interaction and no active heating or cooling demand."
    )
    d.add_paragraph("REQ-DEMO-002 Battery backup runtime")
    d.add_paragraph(
        "Description: On mains power loss, the thermostat shall maintain display and "
        "sensor operation for at least 4 hours from the internal battery."
    )
    d.add_paragraph("REQ-DEMO-003 Charge indicator")
    d.add_paragraph(
        "Description: The thermostat shall display a battery charge indicator whenever "
        "charge is below 20%."
    )

    d.add_picture(str(HERE / "diagram.png"), width=Inches(4))
    d.add_paragraph("Figure 1: System Overview")

    d.add_heading("3.2 User Interface", level=2)
    d.add_paragraph("REQ-DEMO-004 Setpoint adjustment")
    d.add_paragraph(
        "Description: The user shall be able to adjust the target temperature in 0.5 "
        "degree increments using the front-panel dial."
    )
    d.add_paragraph("REQ-DEMO-005 Display readability")
    d.add_paragraph(
        "Description: The display shall remain readable under direct sunlight of up to "
        "10,000 lux."
    )

    d.save(str(OUT))


if __name__ == "__main__":
    build()
    print(f"Wrote {OUT}")
